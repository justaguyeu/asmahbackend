import io
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, A5
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import BusinessSettings

LIME = colors.HexColor("#8CC63F")
INK = colors.HexColor("#111111")
MUTED = colors.HexColor("#8A8A8A")
LINE = colors.HexColor("#E6E6E6")
SOFT = colors.HexColor("#F4F6EF")


def fmt(value, symbol=None):
    symbol = symbol or BusinessSettings.load().currency_symbol
    value = Decimal(value or 0)
    sign = "-" if value < 0 else ""
    return f"{sign}{symbol} {abs(value):,.0f}"


def _styles():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("t", parent=base["Title"], fontName="Helvetica-Bold", fontSize=18,
                                textColor=INK, alignment=0, spaceAfter=2),
        "sub": ParagraphStyle("s", parent=base["Normal"], fontSize=9, textColor=MUTED),
        "h": ParagraphStyle("h", parent=base["Heading3"], fontName="Helvetica-Bold", fontSize=11,
                            textColor=INK, spaceBefore=10, spaceAfter=4),
        "body": ParagraphStyle("b", parent=base["Normal"], fontSize=9, textColor=INK, leading=13),
    }


def _header(story, title, subtitle):
    s = BusinessSettings.load()
    st = _styles()
    logo_cell = ""
    if s.logo:
        try:
            # Read through the storage so this works whether the logo is on disk or in the cloud.
            with s.logo.open("rb") as fh:
                logo_cell = Image(io.BytesIO(fh.read()), width=16 * mm, height=16 * mm, kind="proportional")
        except Exception:
            logo_cell = ""
    contact = " · ".join(x for x in [s.address, s.phone, s.email] if x)
    head = Table([[logo_cell, [Paragraph(s.business_name, st["title"]), Paragraph(contact, st["sub"])]]],
                 colWidths=[20 * mm if logo_cell else 1, None])
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story += [head, Spacer(1, 4),
              Table([[""]], colWidths=["100%"], rowHeights=[2],
                    style=[("BACKGROUND", (0, 0), (-1, -1), LIME)]),
              Spacer(1, 8), Paragraph(title, st["h"]), Paragraph(subtitle, st["sub"]), Spacer(1, 6)]


def _table(rows, widths=None, money_cols=(), header=True):
    t = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
    style = [
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("TEXTCOLOR", (0, 0), (-1, -1), INK),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, LINE),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    if header:
        style += [("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("BACKGROUND", (0, 0), (-1, 0), SOFT),
                  ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#4A4A4A"))]
    for c in money_cols:
        style.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
    t.setStyle(TableStyle(style))
    return t


def receipt_pdf(payment) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A5, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=f"Receipt {payment.receipt_no}")
    st = _styles()
    story = []
    _header(story, f"Payment receipt · {payment.receipt_no}", f"Date: {payment.date:%d %b %Y}")
    loan = payment.loan
    rows = [
        ["Received from", loan.borrower.full_name],
        ["Phone", loan.borrower.phone or "-"],
        ["Loan", loan.reference],
        ["Amount paid", fmt(payment.amount)],
        ["Method", payment.get_method_display()],
        ["Reference", payment.reference or "-"],
        ["Applied to interest", fmt(payment.interest_portion)],
        ["Applied to principal", fmt(payment.principal_portion)],
        ["Applied to penalty", fmt(payment.penalty_portion)],
        ["Balance after payment", fmt(loan.balance)],
        ["Received by", payment.received_by.display_name if payment.received_by else "-"],
    ]
    t = _table(rows, widths=[45 * mm, None], header=False)
    t.setStyle(TableStyle([("FONTNAME", (0, 3), (1, 3), "Helvetica-Bold"), ("FONTSIZE", (0, 3), (1, 3), 11),
                           ("TEXTCOLOR", (0, 0), (0, -1), MUTED)]))
    story.append(t)
    if payment.is_void:
        story += [Spacer(1, 8), Paragraph(f"<font color='#D9534F'><b>VOID</b> — {payment.void_reason}</font>",
                                          st["body"])]
    story += [Spacer(1, 14), Paragraph("Thank you for your payment.", st["sub"])]
    doc.build(story)
    return buf.getvalue()


def statement_pdf(borrower) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm, title=f"Statement {borrower.full_name}")
    st = _styles()
    story = []
    _header(story, f"Borrower statement · {borrower.full_name}",
            f"Phone: {borrower.phone or '-'} · NIDA: {borrower.national_id or '-'} · Printed {date.today():%d %b %Y}")
    loans = borrower.loans.exclude(status="void").order_by("date_issued")
    rows = [["Loan", "Issued", "Principal", "Interest", "Penalty", "Paid", "Balance", "Status"]]
    for l in loans:
        rows.append([l.reference, f"{l.date_issued:%d %b %Y}", fmt(l.principal), fmt(l.total_interest),
                     fmt(l.total_penalty), fmt(l.total_paid), fmt(l.balance), l.get_status_display()])
    story += [Paragraph("Loans", st["h"]), _table(rows, money_cols=(2, 3, 4, 5, 6))]
    pay_rows = [["Receipt", "Date", "Loan", "Method", "Amount"]]
    for l in loans:
        for p in l.payments.filter(is_void=False).order_by("date"):
            pay_rows.append([p.receipt_no, f"{p.date:%d %b %Y}", l.reference, p.get_method_display(), fmt(p.amount)])
    story += [Paragraph("Payments", st["h"]), _table(pay_rows, money_cols=(4,))]
    total_balance = sum((l.balance for l in loans if l.status not in ("completed",)), Decimal(0))
    story += [Spacer(1, 10), Paragraph(f"<b>Total still owed: {fmt(total_balance)}</b>", st["body"])]
    doc.build(story)
    return buf.getvalue()


SUMMARY_LABELS = [
    ("lent", "Money lent out"), ("collected", "Money collected"),
    ("principal_collected", "Principal collected"), ("interest_collected", "Interest collected"),
    ("fees_collected", "Processing fees"), ("penalties_collected", "Penalties collected"),
    ("profit", "Profit made"), ("expected_profit", "Expected profit"),
    ("outstanding", "Still outstanding (loans from this period)"), ("overdue_amount", "Overdue"),
    ("lost_amount", "Defaulted / written off"),
]
COUNT_LABELS = [("loans_issued", "Loans issued"), ("new_borrowers", "New borrowers"),
                ("collection_rate", "Collection rate (%)")]
MONTH_HEAD = ["Month", "Lent", "Collected", "Interest earned", "Fees/Penalties", "Profit", "Outstanding", "Overdue"]
MONTH_KEYS = ["lent", "collected", "interest_collected", "fees_and_penalties", "profit", "outstanding",
              "overdue_amount"]


def period_xlsx(summary, rows, totals, borrowers) -> bytes:
    s = BusinessSettings.load()
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="EEF5E3")
    money = '#,##0;[Red]-#,##0'
    thin = Border(bottom=Side(style="thin", color="E6E6E6"))

    ws["A1"] = s.business_name
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Report for {summary['start']:%d %b %Y} – {summary['end']:%d %b %Y} ({s.currency})"
    ws["A2"].font = Font(color="8A8A8A")
    r = 4
    for key, label in SUMMARY_LABELS:
        ws.cell(r, 1, label)
        c = ws.cell(r, 2, float(summary[key] or 0))
        c.number_format = money
        if key == "profit":
            ws.cell(r, 1).font = bold
            c.font = bold
        r += 1
    for key, label in COUNT_LABELS:
        ws.cell(r, 1, label)
        ws.cell(r, 2, float(summary[key]) if summary[key] is not None else "-")
        r += 1
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18

    ws2 = wb.create_sheet("Month by month")
    ws2.append(MONTH_HEAD)
    for row in rows:
        ws2.append([row["label"]] + [float(row[k]) for k in MONTH_KEYS])
    ws2.append(["TOTAL"] + [float(totals[k]) for k in MONTH_KEYS])
    for cell in ws2[1]:
        cell.font = bold
        cell.fill = head_fill
    for cell in ws2[ws2.max_row]:
        cell.font = bold
    for col in range(1, len(MONTH_HEAD) + 1):
        ws2.column_dimensions[get_column_letter(col)].width = 16
        for row_cells in ws2.iter_rows(min_row=2, min_col=col, max_col=col):
            for cell in row_cells:
                cell.border = thin
                if col > 1:
                    cell.number_format = money

    ws3 = wb.create_sheet("Borrowers")
    ws3.append(["Name", "Phone", "Status", "Borrowed in period", "Paid in period", "Balance now"])
    for b in borrowers:
        ws3.append([b["full_name"], b["phone"], b["status"].title(), float(b["borrowed"]), float(b["paid"]),
                    float(b["balance"])])
    for cell in ws3[1]:
        cell.font = bold
        cell.fill = head_fill
    for col, width in zip("ABCDEF", [28, 16, 12, 18, 18, 18]):
        ws3.column_dimensions[col].width = width
    for row_cells in ws3.iter_rows(min_row=2, min_col=4, max_col=6):
        for cell in row_cells:
            cell.number_format = money

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def period_pdf(summary, rows, totals, borrowers) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title="ASK Company report")
    st = _styles()
    story = []
    _header(story, "Lending report", f"{summary['start']:%d %b %Y} – {summary['end']:%d %b %Y}")
    cards = [[label, fmt(summary[key])] for key, label in SUMMARY_LABELS]
    cards += [[label, (f"{summary[key]}%" if key == "collection_rate" else str(summary[key]))
               if summary[key] is not None else "-"] for key, label in COUNT_LABELS]
    t = _table(cards, widths=[90 * mm, 50 * mm], money_cols=(1,), header=False)
    idx = [k for k, _ in SUMMARY_LABELS].index("profit")
    t.setStyle(TableStyle([("FONTNAME", (0, idx), (1, idx), "Helvetica-Bold"),
                           ("BACKGROUND", (0, idx), (1, idx), SOFT)]))
    story += [t, Paragraph("Month by month", st["h"])]
    table_rows = [MONTH_HEAD] + [[r["label"]] + [fmt(r[k], "") for k in MONTH_KEYS] for r in rows]
    table_rows.append(["TOTAL"] + [fmt(totals[k], "") for k in MONTH_KEYS])
    mt = _table(table_rows, money_cols=range(1, 8))
    mt.setStyle(TableStyle([("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7.5)]))
    story.append(mt)
    if borrowers:
        story.append(Paragraph("Borrowers active in this period", st["h"]))
        b_rows = [["Name", "Phone", "Status", "Borrowed", "Paid", "Balance"]]
        b_rows += [[b["full_name"], b["phone"] or "-", b["status"].title(), fmt(b["borrowed"], ""),
                    fmt(b["paid"], ""), fmt(b["balance"], "")] for b in borrowers]
        story.append(_table(b_rows, money_cols=(3, 4, 5)))
    doc.build(story)
    return buf.getvalue()
