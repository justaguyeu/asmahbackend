import csv
import io
from datetime import date, datetime

from django.core import management
from django.db import transaction
from django.db.models import Count, DecimalField, Q, Sum
from django.db.models.functions import Coalesce
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import mixins, status, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import ValidationError
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.response import Response
from rest_framework.views import APIView

from . import exports, reports, services
from .activity import log_activity
from .models import ZERO, ActivityLog, Borrower, BusinessSettings, Loan, Payment, with_status
from .serializers import (
    ActivityLogSerializer,
    BorrowerSerializer,
    BulkLoanSerializer,
    BusinessSettingsSerializer,
    LoanCreateSerializer,
    LoanDetailSerializer,
    LoanListSerializer,
    LoanPreviewSerializer,
    PaymentSerializer,
)

DEC = DecimalField(max_digits=16, decimal_places=2)
OPEN = reports.OPEN_STATUSES


def file_response(content, filename, content_type):
    response = HttpResponse(content, content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response["Access-Control-Expose-Headers"] = "Content-Disposition"
    return response


def parse_date(value, field):
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        raise ValidationError({field: "Use the format YYYY-MM-DD."})


def period_from_request(request):
    today = timezone.localdate()
    start = parse_date(request.query_params.get("start"), "start") or today.replace(month=1, day=1)
    end = parse_date(request.query_params.get("end"), "end") or today
    if end < start:
        raise ValidationError({"end": "End date must be after the start date."})
    if (end - start).days > 366 * 15:
        raise ValidationError({"start": "Please choose a period shorter than 15 years."})
    return start, end


class BorrowerViewSet(mixins.ListModelMixin, mixins.RetrieveModelMixin, mixins.CreateModelMixin,
                      mixins.UpdateModelMixin, viewsets.GenericViewSet):
    serializer_class = BorrowerSerializer
    parser_classes = [JSONParser, MultiPartParser, FormParser]

    def get_queryset(self):
        live = Q(loans__status__in=OPEN)
        qs = with_status(Borrower.objects).annotate(
            balance=Coalesce(Sum("loans__balance", filter=live), ZERO, output_field=DEC),
            total_borrowed=Coalesce(Sum("loans__principal", filter=~Q(loans__status="void")), ZERO,
                                    output_field=DEC),
            active_loans=Count("loans", filter=live & Q(loans__balance__gt=0)),
            loans_count=Count("loans", filter=~Q(loans__status="void")),
        )
        p = self.request.query_params
        if self.action == "list":
            qs = qs.filter(is_active=p.get("archived") != "1")
        if q := p.get("q", "").strip():
            qs = qs.filter(Q(full_name__icontains=q) | Q(phone__icontains=q) | Q(national_id__icontains=q))
        status_filter = p.get("status")
        if status_filter == "defaulted":
            qs = qs.filter(status_lost=True)
        elif status_filter == "late":
            qs = qs.filter(status_late=True, status_lost=False)
        elif status_filter == "good":
            qs = qs.filter(status_late=False, status_lost=False)
        elif status_filter == "active":
            qs = qs.filter(active_loans__gt=0)
        ordering = {"name": "full_name", "-balance": "-balance", "balance": "balance",
                    "-created": "-created_at"}.get(p.get("ordering", "name"), "full_name")
        return qs.order_by(ordering, "id")

    def perform_create(self, serializer):
        b = serializer.save(created_by=self.request.user)
        log_activity(self.request.user, "borrower_created", "borrower", b.pk, f"Added borrower {b.full_name}")

    def perform_update(self, serializer):
        b = serializer.save()
        log_activity(self.request.user, "borrower_updated", "borrower", b.pk, f"Updated borrower {b.full_name}")

    @action(detail=True, methods=["post"])
    def archive(self, request, pk=None):
        b = self.get_object()
        b.is_active = request.data.get("restore") is not True
        b.save(update_fields=["is_active"])
        verb = "Archived" if not b.is_active else "Restored"
        log_activity(request.user, "borrower_archived", "borrower", b.pk, f"{verb} borrower {b.full_name}")
        return Response(self.get_serializer(self.get_queryset().get(pk=b.pk)).data)

    @action(detail=False, methods=["post"])
    def bulk(self, request):
        """Add many borrowers at once. Accepts {"text": "Name, phone\\nName, phone"} or
        {"rows": [{full_name, phone, ...}]}."""
        rows = request.data.get("rows")
        if rows is None:
            rows = []
            for line in (request.data.get("text") or "").splitlines():
                parts = [x.strip() for x in line.split(",")]
                if parts and parts[0]:
                    rows.append({"full_name": parts[0], "phone": parts[1] if len(parts) > 1 else "",
                                 "address": parts[2] if len(parts) > 2 else ""})
        return self._create_many(rows, "bulk add")

    @action(detail=False, methods=["post"], url_path="import")
    def import_file(self, request):
        upload = request.FILES.get("file")
        if not upload:
            raise ValidationError({"file": "Choose a CSV or Excel file."})
        name = upload.name.lower()
        rows = []
        if name.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            ws = load_workbook(upload, read_only=True, data_only=True).active
            data = [[("" if c is None else str(c)).strip() for c in r] for r in ws.iter_rows(values_only=True)]
        elif name.endswith(".csv"):
            text = upload.read().decode("utf-8-sig", errors="replace")
            data = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text))]
        else:
            raise ValidationError({"file": "Only .csv and .xlsx files are supported."})
        if not data:
            raise ValidationError({"file": "The file is empty."})
        header = [h.lower().replace(" ", "_") for h in data[0]]
        aliases = {"name": "full_name", "full_name": "full_name", "jina": "full_name",
                   "phone": "phone", "phone_number": "phone", "simu": "phone",
                   "national_id": "national_id", "nida": "national_id", "id_number": "national_id",
                   "address": "address", "location": "address", "guarantor_name": "guarantor_name",
                   "guarantor": "guarantor_name", "guarantor_phone": "guarantor_phone", "notes": "notes"}
        cols = {i: aliases[h] for i, h in enumerate(header) if h in aliases}
        if "full_name" not in cols.values():
            cols, body = {0: "full_name", 1: "phone", 2: "address"}, data
        else:
            body = data[1:]
        for r in body:
            row = {field: r[i] for i, field in cols.items() if i < len(r)}
            if row.get("full_name"):
                rows.append(row)
        return self._create_many(rows, f"import from {upload.name}")

    def _create_many(self, rows, source):
        created, skipped, errors = [], 0, []
        existing = {(b.full_name.lower(), b.phone) for b in Borrower.objects.all()}
        with transaction.atomic():
            for i, row in enumerate(rows, start=1):
                ser = BorrowerSerializer(data=row)
                if not ser.is_valid():
                    errors.append({"row": i, "errors": ser.errors})
                    continue
                key = (ser.validated_data["full_name"].lower(), ser.validated_data.get("phone", ""))
                if key in existing:
                    skipped += 1
                    continue
                existing.add(key)
                created.append(ser.save(created_by=self.request.user))
        if created:
            log_activity(self.request.user, "borrowers_bulk", "borrower", "",
                         f"Added {len(created)} borrowers by {source}")
        return Response({"created": len(created), "skipped_duplicates": skipped, "errors": errors[:20]},
                        status=status.HTTP_201_CREATED if created else status.HTTP_200_OK)

    @action(detail=False, methods=["get"], url_path="import-template")
    def import_template(self, request):
        content = "full_name,phone,national_id,address,guarantor_name,guarantor_phone,notes\n" \
                  "Amina Juma,0712345678,19900101-12345-00001-12,Kariakoo,Hassan Juma,0754000111,\n"
        return file_response(content, "borrowers-template.csv", "text/csv")

    @action(detail=True, methods=["get"])
    def statement(self, request, pk=None):
        b = self.get_object()
        name = b.full_name.replace(" ", "-").lower()
        return file_response(exports.statement_pdf(b), f"statement-{name}.pdf", "application/pdf")

    @action(detail=True, methods=["get"])
    def history(self, request, pk=None):
        b = self.get_object()
        loans = b.loans.order_by("-date_issued", "-id")
        pays = Payment.objects.filter(loan__borrower=b).select_related("loan", "received_by", "voided_by")
        return Response({"loans": LoanListSerializer(loans, many=True).data,
                         "payments": PaymentSerializer(pays, many=True).data})


class LoanViewSet(mixins.ListModelMixin, mixins.RetrieveModelMixin, viewsets.GenericViewSet):
    def get_serializer_class(self):
        return LoanDetailSerializer if self.action == "retrieve" else LoanListSerializer

    def get_queryset(self):
        qs = Loan.objects.select_related("borrower", "created_by")
        if self.action == "retrieve":
            return qs.prefetch_related("installments", "payments__received_by", "payments__voided_by",
                                       "payments__loan__borrower")
        p = self.request.query_params
        status_filter = p.get("status")
        if status_filter == "open":
            qs = qs.filter(status__in=OPEN)
        elif status_filter == "overdue":
            qs = qs.filter(status__in=[Loan.Status.OVERDUE, Loan.Status.DEFAULTED])
        elif status_filter:
            qs = qs.filter(status=status_filter)
        else:
            qs = qs.exclude(status=Loan.Status.VOID)
        if borrower := p.get("borrower"):
            qs = qs.filter(borrower_id=borrower)
        if q := p.get("q", "").strip():
            qs = qs.filter(Q(reference__icontains=q) | Q(borrower__full_name__icontains=q)
                           | Q(borrower__phone__icontains=q))
        if start := parse_date(p.get("start"), "start"):
            qs = qs.filter(date_issued__gte=start)
        if end := parse_date(p.get("end"), "end"):
            qs = qs.filter(date_issued__lte=end)
        ordering = {"-balance": "-balance", "balance": "balance", "date": "date_issued",
                    "-date": "-date_issued", "-principal": "-principal"}.get(p.get("ordering"), "-date_issued")
        return qs.order_by(ordering, "-id")

    def list(self, request, *args, **kwargs):
        services.refresh_open_loans()
        return super().list(request, *args, **kwargs)

    def retrieve(self, request, *args, **kwargs):
        loan = self.get_object()
        # every change to a loan recalculates it; this only catches up with the passing days
        if loan.status not in (Loan.Status.VOID, Loan.Status.COMPLETED) and loan.checked_on != timezone.localdate():
            services.recalculate_loan(loan)
        return super().retrieve(request, *args, **kwargs)

    def create(self, request):
        ser = LoanCreateSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        data = dict(ser.validated_data)
        borrower = data.pop("borrower")
        confirm = data.pop("confirm_existing", False)
        unpaid = borrower.loans.filter(status__in=OPEN, balance__gt=0)
        if unpaid.exists() and not confirm:
            total = unpaid.aggregate(s=Sum("balance"))["s"]
            return Response({
                "code": "existing_loan",
                "detail": f"{borrower.full_name} still owes {BusinessSettings.load().currency_symbol} "
                          f"{total:,.0f} on {unpaid.count()} unpaid loan(s). Give another loan anyway?",
            }, status=status.HTTP_409_CONFLICT)
        loan = services.create_loan(borrower=borrower, user=request.user, **data)
        log_activity(request.user, "loan_created", "loan", loan.pk,
                     f"Gave loan {loan.reference} of {loan.principal:,.0f} to {borrower.full_name}")
        loan = self.get_queryset().model.objects.get(pk=loan.pk)
        return Response(LoanDetailSerializer(loan).data, status=status.HTTP_201_CREATED)

    def partial_update(self, request, pk=None):
        loan = self.get_object()
        if "notes" in request.data:
            loan.notes = str(request.data["notes"])[:2000]
            loan.save(update_fields=["notes"])
            log_activity(request.user, "loan_updated", "loan", loan.pk, f"Updated notes on {loan.reference}")
        return Response(LoanDetailSerializer(loan).data)

    @action(detail=False, methods=["post"])
    def preview(self, request):
        """Calculate a repayment schedule without saving anything."""
        ser = LoanPreviewSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        d = ser.validated_data
        s = BusinessSettings.load()
        rate = d.get("interest_rate", s.default_interest_rate)
        itype = d.get("interest_type") or s.default_interest_type
        months = d.get("duration_months") or s.default_duration_months
        freq = d.get("frequency") or s.default_frequency
        issued = d.get("date_issued") or timezone.localdate()
        first = d.get("first_due_date") or services.default_first_due(issued, freq)
        fee = d.get("processing_fee")
        if fee is None:
            fee = services.compute_fee(d["principal"], s.processing_fee_type, s.processing_fee_value)
        rows = services.build_schedule(d["principal"], rate, itype, months, freq, first)
        total_interest = sum((r["interest_due"] for r in rows), ZERO)
        principal = services.to_money(d["principal"])
        return Response({
            "schedule": [{**r, "amount": r["principal_due"] + r["interest_due"]} for r in rows],
            "total_interest": total_interest, "total_due": principal + total_interest,
            "processing_fee": services.to_money(fee), "disbursed_amount": principal - services.to_money(fee),
            "installment_amount": rows[0]["principal_due"] + rows[0]["interest_due"],
            "count": len(rows), "interest_rate": rate, "interest_type": itype,
            "duration_months": months, "frequency": freq, "first_due_date": first,
        })

    @action(detail=False, methods=["post"])
    def bulk(self, request):
        """Monthly lending round: same terms, many borrowers, one click."""
        ser = BulkLoanSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        d = dict(ser.validated_data)
        items = d.pop("items")
        created = []
        with transaction.atomic():
            for item in items:
                loan = services.create_loan(borrower=item["borrower"], principal=item["principal"],
                                            user=request.user, **d)
                created.append(loan)
        total = sum((l.principal for l in created), ZERO)
        log_activity(request.user, "loan_round", "loan", "",
                     f"Monthly round: {len(created)} loans, total {total:,.0f}")
        return Response({"created": len(created), "total": total,
                         "loans": LoanListSerializer(created, many=True).data}, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=["post"])
    def void(self, request, pk=None):
        loan = self.get_object()
        reason = (request.data.get("reason") or "").strip()
        if len(reason) < 3:
            raise ValidationError({"reason": "Give a reason for voiding this loan."})
        if loan.status == Loan.Status.VOID:
            raise ValidationError({"detail": "This loan is already void."})
        services.void_loan(loan, reason, request.user)
        log_activity(request.user, "loan_voided", "loan", loan.pk, f"Voided loan {loan.reference}: {reason}")
        return Response(LoanDetailSerializer(Loan.objects.get(pk=loan.pk)).data)

    @action(detail=True, methods=["post"], url_path="set-status")
    def set_status(self, request, pk=None):
        loan = self.get_object()
        new = request.data.get("status")
        if new not in ("defaulted", "written_off", "auto"):
            raise ValidationError({"status": "Choose defaulted, written_off or auto."})
        if loan.status in (Loan.Status.VOID, Loan.Status.COMPLETED):
            raise ValidationError({"status": "This loan is closed."})
        services.set_loan_status(loan, new)
        log_activity(request.user, "loan_status", "loan", loan.pk, f"Set {loan.reference} status to {new}")
        return Response(LoanDetailSerializer(Loan.objects.get(pk=loan.pk)).data)


class PaymentViewSet(mixins.ListModelMixin, mixins.RetrieveModelMixin, mixins.CreateModelMixin,
                     viewsets.GenericViewSet):
    serializer_class = PaymentSerializer

    def get_queryset(self):
        qs = Payment.objects.select_related("loan__borrower", "received_by", "voided_by")
        p = self.request.query_params
        if p.get("void") == "1":
            qs = qs.filter(is_void=True)
        elif p.get("void") != "all":
            qs = qs.filter(is_void=False)
        if loan := p.get("loan"):
            qs = qs.filter(loan_id=loan)
        if method := p.get("method"):
            qs = qs.filter(method=method)
        if q := p.get("q", "").strip():
            qs = qs.filter(Q(receipt_no__icontains=q) | Q(reference__icontains=q)
                           | Q(loan__borrower__full_name__icontains=q) | Q(loan__reference__icontains=q))
        if start := parse_date(p.get("start"), "start"):
            qs = qs.filter(date__gte=start)
        if end := parse_date(p.get("end"), "end"):
            qs = qs.filter(date__lte=end)
        return qs.order_by("-date", "-id")

    def list(self, request, *args, **kwargs):
        response = super().list(request, *args, **kwargs)
        totals = self.filter_queryset(self.get_queryset()).aggregate(
            total=Coalesce(Sum("amount"), ZERO, output_field=DEC),
            interest=Coalesce(Sum("interest_portion"), ZERO, output_field=DEC))
        response.data["totals"] = totals
        return response

    def create(self, request, *args, **kwargs):
        ser = self.get_serializer(data=request.data)
        ser.is_valid(raise_exception=True)
        d = ser.validated_data
        loan = d["loan"]
        if loan.balance <= 0:
            raise ValidationError({"loan": "This loan is already fully paid."})
        payment = services.record_payment(loan=loan, amount=d["amount"], date_paid=d.get("date"),
                                          method=d.get("method", "cash"), reference=d.get("reference", ""),
                                          notes=d.get("notes", ""), user=request.user)
        log_activity(request.user, "payment_recorded", "payment", payment.pk,
                     f"Recorded {payment.amount:,.0f} from {loan.borrower.full_name} on {loan.reference} "
                     f"({payment.receipt_no})")
        return Response(PaymentSerializer(payment).data, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=["post"])
    def void(self, request, pk=None):
        payment = self.get_object()
        reason = (request.data.get("reason") or "").strip()
        if len(reason) < 3:
            raise ValidationError({"reason": "Give a reason for voiding this payment."})
        if payment.is_void:
            raise ValidationError({"detail": "Already void."})
        services.void_payment(payment, reason, request.user)
        log_activity(request.user, "payment_voided", "payment", payment.pk,
                     f"Voided payment {payment.receipt_no} ({payment.amount:,.0f}): {reason}")
        return Response(PaymentSerializer(payment).data)

    @action(detail=True, methods=["get"])
    def receipt(self, request, pk=None):
        payment = self.get_object()
        return file_response(exports.receipt_pdf(payment), f"receipt-{payment.receipt_no}.pdf", "application/pdf")


class SettingsView(APIView):
    parser_classes = [JSONParser, MultiPartParser, FormParser]

    def get(self, request):
        return Response(BusinessSettingsSerializer(BusinessSettings.load(), context={"request": request}).data)

    def patch(self, request):
        ser = BusinessSettingsSerializer(BusinessSettings.load(), data=request.data, partial=True,
                                         context={"request": request})
        ser.is_valid(raise_exception=True)
        ser.save()
        log_activity(request.user, "settings_updated", "settings", 1,
                     "Updated settings: " + ", ".join(sorted(request.data.keys()))[:400])
        return Response(ser.data)


class DashboardView(APIView):
    def get(self, request):
        services.refresh_open_loans()
        range_key = request.query_params.get("range", "7d")
        if range_key not in ("7d", "30d", "12m"):
            range_key = "7d"
        return Response(reports.dashboard_data(range_key))


class AlertsView(APIView):
    def get(self, request):
        services.refresh_open_loans()
        items = reports.overdue_items()
        return Response({"overdue_count": len(items), "items": items[:8]})


class ReportSummaryView(APIView):

    def get(self, request):
        services.refresh_open_loans()
        start, end = period_from_request(request)
        rows, totals = reports.monthly_rows(start, end)
        return Response({"summary": reports.period_summary(start, end), "months": rows, "totals": totals,
                         "years": reports.yearly_rows()})


class ReportCompareView(APIView):

    def get(self, request):
        p = request.query_params
        a = (parse_date(p.get("a_start"), "a_start"), parse_date(p.get("a_end"), "a_end"))
        b = (parse_date(p.get("b_start"), "b_start"), parse_date(p.get("b_end"), "b_end"))
        if not all(a + b):
            raise ValidationError({"detail": "Choose both periods."})
        return Response({"a": reports.period_summary(*a), "b": reports.period_summary(*b)})


class ReportMonthView(APIView):

    def get(self, request):
        try:
            year, month = int(request.query_params["year"]), int(request.query_params["month"])
            date(year, month, 1)
        except (KeyError, ValueError):
            raise ValidationError({"detail": "Give a valid year and month."})
        return Response(reports.month_detail(year, month))


class ReportBorrowersView(APIView):

    def get(self, request):
        start, end = period_from_request(request)
        return Response(reports.borrowers_in_period(start, end))


class ReportExportView(APIView):

    def get(self, request):
        start, end = period_from_request(request)
        fmt = request.query_params.get("type", "xlsx")
        summary = reports.period_summary(start, end)
        rows, totals = reports.monthly_rows(start, end)
        borrowers = reports.borrowers_in_period(start, end)
        name = f"ask-company-report-{start:%Y%m%d}-{end:%Y%m%d}"
        log_activity(request.user, "report_exported", "report", "", f"Exported {fmt.upper()} report {start} to {end}")
        if fmt == "pdf":
            return file_response(exports.period_pdf(summary, rows, totals, borrowers), f"{name}.pdf",
                                 "application/pdf")
        return file_response(exports.period_xlsx(summary, rows, totals, borrowers), f"{name}.xlsx",
                             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


class ActivityViewSet(mixins.ListModelMixin, viewsets.GenericViewSet):
    serializer_class = ActivityLogSerializer

    def get_queryset(self):
        qs = ActivityLog.objects.select_related("user")
        p = self.request.query_params
        if q := p.get("q", "").strip():
            qs = qs.filter(Q(description__icontains=q) | Q(user__username__icontains=q))
        if action_name := p.get("action"):
            qs = qs.filter(action=action_name)
        if user := p.get("user"):
            qs = qs.filter(user_id=user)
        return qs


class BackupView(APIView):
    """Download a full JSON backup of the database (admin only)."""

    def get(self, request):
        out = io.StringIO()
        management.call_command("dumpdata", "accounts", "lending", indent=1, stdout=out,
                                exclude=["accounts.user"] if request.query_params.get("no_users") else [])
        log_activity(request.user, "backup_downloaded", "backup", "", "Downloaded a database backup")
        stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
        return file_response(out.getvalue(), f"ask-company-backup-{stamp}.json", "application/json")


class HealthView(APIView):
    permission_classes = []
    authentication_classes = []

    def get(self, request):
        return Response({"ok": True, "time": timezone.now()})

